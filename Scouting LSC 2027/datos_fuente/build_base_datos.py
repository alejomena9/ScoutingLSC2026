import csv, re, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
src='datos_fuente/usl_championship_2026_fotmob.csv'
R=list(csv.DictReader(open(src,encoding='utf-8')))
conf={r['teamId']:r['conf'] for r in R if r['conf']}
tname={r['teamId']:r['team'] for r in R if r['conf']}
grp={'keepers':'Portero','defenders':'Defensa','midfielders':'Mediocampista','attackers':'Delantero','NOT_IN_CURRENT_SQUAD':''}
def slug(s): s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower(); return re.sub(r'[^a-z0-9]+','-',s).strip('-')
def num(v):
    if v in ('',None): return None
    try: f=float(v); return int(f) if f.is_integer() else f
    except: return v
cols=[('id','ID FotMob'),('name','Jugador'),('team','Club'),('conf','Conferencia'),('estado','Estado'),('linea','Línea'),('pos1','Pos. principal'),('posDesc','Posiciones FotMob'),('shirt','Dorsal'),('age','Edad'),('dob','Nacimiento'),('nat','Nacionalidad'),('height','Altura (cm)'),('value','Valor est. FotMob (€)'),('injury','Lesionado'),
('mp','PJ'),('mins_played','Minutos'),('rating','Rating FotMob'),('potm','MVP partido'),('goals','Goles'),('pen_goals','Goles de penal'),('goal_assist','Asistencias'),('_goals_and_goal_assist','G+A'),('goals_per_90','Goles /90'),
('total_scoring_att','Tiros /90'),('ontarget_scoring_att','Tiros a puerta /90'),('shot_conv','Conversión tiro %'),('big_chance_missed','Ocasiones claras falladas'),
('total_att_assist','Ocasiones creadas'),('chances_p90','Ocasiones creadas /90'),('big_chance_created','Ocasiones claras creadas'),('accurate_pass','Pases precisos /90'),('pass_pct','Precisión pase %'),('accurate_long_balls','Balones largos precisos /90'),('longball_pct','Balón largo %'),
('won_contest','Regates exitosos /90'),('dribble_pct','Regate %'),('penalty_won','Penales provocados'),('fouled_p90','Faltas recibidas /90'),
('defensive_contributions','Acciones defensivas /90'),('total_tackle','Entradas /90'),('interception','Intercepciones /90'),('effective_clearance','Despejes /90'),('outfielder_block','Bloqueos /90'),('ball_recovery','Recuperaciones /90'),('poss_won_att_3rd','Recup. último tercio /90'),('poss_won_mid_p90','Recup. medio campo /90'),
('fouls','Faltas cometidas /90'),('yellow_card','Amarillas'),('red_card','Rojas'),('penalty_conceded','Penales cometidos'),
('clean_sheet','Porterías a cero'),('saves','Atajadas /90'),('_save_percentage','% Atajadas'),('goals_conceded','Goles recibidos /90'),('url','Enlace FotMob')]
rows=[]
for r in R:
    d=dict(r); out=r['group']=='NOT_IN_CURRENT_SQUAD'
    d['conf']=r['conf'] or conf.get(r['teamId'],'')
    if out: d['team']=tname.get(r['teamId'], r['team'])
    d['estado']='Fuera de plantilla actual' if out else 'En plantilla'
    d['linea']=grp[r['group']] if not out else ''
    d['pos1']=r['posDesc'].split(',')[0] if r['posDesc'] else ''
    if out and not d['linea'] and r['statPos']:
        p=int(r['statPos'].split('|')[0]); d['linea']='Portero' if p==11 else 'Defensa' if p<60 else 'Mediocampista' if p<100 else 'Delantero'
    d['injury']='Sí' if r['injury'] else ''
    d['url']=f"https://www.fotmob.com/players/{r['id']}/{slug(r['name'])}"
    rows.append(d)
order={'Portero':0,'Defensa':1,'Mediocampista':2,'Delantero':3,'':4}
rows.sort(key=lambda d:(d['estado']!='En plantilla', d['team'], order[d['linea']], -(num(d['mins_played']) or 0)))
wb=Workbook(); ws=wb.active; ws.title='Jugadores'
F='Arial'; hdr=PatternFill('solid',fgColor='111111'); green=PatternFill('solid',fgColor='3DDC5A')
ws.append([c[1] for c in cols])
for d in rows: ws.append([num(d.get(k,'')) if k not in ('name','team','dob','nat','posDesc','url','conf','estado','linea','pos1','injury') else (d.get(k) or None) for k,_ in cols])
for c in ws[1]: c.font=Font(name=F,bold=True,color='F5F1E9'); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
ws.row_dimensions[1].height=42
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=Font(name=F,size=10)
ci={k:i+1 for i,(k,_) in enumerate(cols)}
for row in ws.iter_rows(min_row=2):
    row[ci['value']-1].number_format='#,##0'; row[ci['mins_played']-1].number_format='#,##0'
    u=row[ci['url']-1]; u.hyperlink=u.value; u.font=Font(name=F,size=10,color='1F6F2E',underline='single')
widths={'name':24,'team':26,'posDesc':18,'nat':16,'url':44,'estado':20,'dob':11,'value':13}
for k,i in ci.items(): ws.column_dimensions[get_column_letter(i)].width=widths.get(k,11)
ws.freeze_panes='C2'
t=Table(displayName='USL2026',ref=f"A1:{get_column_letter(len(cols))}{len(rows)+1}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=True); ws.add_table(t)
# Equipos
we=wb.create_sheet('Equipos')
we.append(['Club','Conferencia','Jugadores en plantilla','Porteros','Defensas','Mediocampistas','Delanteros','Edad media','Jugadores fuera de plantilla (jugaron 2026)'])
teams=sorted({(d['team'],d['conf']) for d in rows if d['estado']=='En plantilla'},key=lambda x:(x[1],x[0]))
for i,(tn,cf) in enumerate(teams,start=2):
    J="Jugadores"; 
    we.append([tn,cf,f'=COUNTIFS({J}!$C:$C,A{i},{J}!$E:$E,"En plantilla")',f'=COUNTIFS({J}!$C:$C,A{i},{J}!$F:$F,"Portero")',f'=COUNTIFS({J}!$C:$C,A{i},{J}!$F:$F,"Defensa")',f'=COUNTIFS({J}!$C:$C,A{i},{J}!$F:$F,"Mediocampista")',f'=COUNTIFS({J}!$C:$C,A{i},{J}!$F:$F,"Delantero")',f'=AVERAGEIFS({J}!$J:$J,{J}!$C:$C,A{i},{J}!$E:$E,"En plantilla")',f'=COUNTIFS({J}!$C:$C,A{i},{J}!$E:$E,"Fuera de plantilla actual")'])
n=len(teams)+1
we.append(['TOTAL','',f'=SUM(C2:C{n})',f'=SUM(D2:D{n})',f'=SUM(E2:E{n})',f'=SUM(F2:F{n})',f'=SUM(G2:G{n})',f'=AVERAGEIFS(Jugadores!$J:$J,Jugadores!$E:$E,"En plantilla")',f'=SUM(I2:I{n})'])
for c in we[1]: c.font=Font(name=F,bold=True,color='F5F1E9'); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='center')
for row in we.iter_rows(min_row=2):
    for c in row: c.font=Font(name=F,size=10,bold=(row[0].value=='TOTAL'))
    row[7].number_format='0.0'
for c in we[n+1]: c.fill=green
for i,w in enumerate([30,13,13,10,10,14,11,11,22],1): we.column_dimensions[get_column_letter(i)].width=w
we.freeze_panes='A2'
# Diccionario
wd=wb.create_sheet('Notas')
notes=[('Fuente','FotMob — USL Championship 2026 (liga 8972, temporada 30301). Plantillas: fotmob.com/teams/[id]/squad. Estadísticas: data.fotmob.com/stats/8972/season/30301/*.json'),
('Fecha de extracción','30 sep 2026 (tras ~26 jornadas de temporada regular)'),
('Cobertura',f'{sum(1 for d in rows if d["estado"]=="En plantilla")} jugadores en plantillas actuales de los 25 clubes + {sum(1 for d in rows if d["estado"]!="En plantilla")} que jugaron en 2026 pero ya no figuran en ninguna plantilla (cedidos/traspasados/liberados). Entrenadores excluidos.'),
('Métricas /90','FotMob solo publica las métricas por 90 min a partir de un mínimo de minutos; una celda vacía significa "sin dato / por debajo del umbral", no cero.'),
('Valor est. FotMob (€)','Estimación algorítmica de FotMob, no valor de mercado real ni salario. Útil solo como referencia relativa.'),
('Lexington SC','LSC aparece incluido como referencia de comparación; filtrar Club ≠ "Lexington SC" para ver objetivos externos.'),
('Rating FotMob / MVP partido','Rating medio de temporada; MVP = veces nombrado jugador del partido.'),
('Estado del contrato','No disponible en FotMob — pendiente de otra fuente si se usa como criterio.'),
('Próximo paso','Añadir criterios de scouting (posición, edad, minutos mínimos, métricas clave) y construir ranking/percentiles por posición.')]
wd.append(['Campo','Detalle'])
for n_ in notes: wd.append(list(n_))
for c in wd[1]: c.font=Font(name=F,bold=True,color='F5F1E9'); c.fill=hdr
for row in wd.iter_rows(min_row=2):
    row[0].font=Font(name=F,bold=True,size=10); row[1].font=Font(name=F,size=10); row[1].alignment=Alignment(wrap_text=True,vertical='top'); row[0].alignment=Alignment(vertical='top')
wd.column_dimensions['A'].width=26; wd.column_dimensions['B'].width=110
wb.save('USL_Championship_2026_BaseDatos_Jugadores.xlsx'); print('ok',len(rows))
